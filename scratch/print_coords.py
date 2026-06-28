from pptx import Presentation
prs = Presentation(r'd:\PS15_SolarFlare\BAH_PPT (1).pptx')

print('Slide Width:', prs.slide_width / 914400)
print('Slide Height:', prs.slide_height / 914400)

print('\n=== Slide 7 Shapes ===')
slide7 = prs.slides[6]
for s_idx, shape in enumerate(slide7.shapes):
    left = shape.left / 914400
    top = shape.top / 914400
    w = shape.width / 914400
    h = shape.height / 914400
    print(f'  Shape {s_idx}: {shape.name} | Type: {shape.shape_type} | Left: {left:.2f}\" | Top: {top:.2f}\" | Width: {w:.2f}\" | Height: {h:.2f}\"')
